import csv
from pathlib import Path

from openpyxl import load_workbook

from lead_pipeline.export import write_csv, write_xlsx
from lead_pipeline.llm import FakeProvider
from lead_pipeline.pipeline import process
from lead_pipeline.schema import ROW_COLUMNS, LeadRecord


def test_csv_and_xlsx_exports(tmp_path: Path, records: list[LeadRecord], prompt) -> None:
    results = process(records, FakeProvider(fail_ids={"T-5"}), prompt, workers=1)

    csv_path = write_csv(results, tmp_path / "out" / "leads.csv")
    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 5
    assert list(rows[0].keys()) == list(ROW_COLUMNS)
    assert rows[0]["category"] == "запись"
    assert rows[0]["phone"] == "+79130000021"
    assert rows[4]["error"].startswith("искусственная ошибка")

    xlsx_path = write_xlsx(results, tmp_path / "out" / "leads.xlsx")
    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == ["Заявки", "Сводка"]
    ws = wb["Заявки"]
    assert ws.max_row == 6
    assert ws["D2"].value == "запись"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    summary = {row[0]: row[1] for row in wb["Сводка"].iter_rows(values_only=True) if row[0]}
    assert summary["Всего заявок"] == 5
    assert summary["Ошибок"] == 1
    assert summary["жалоба"] == 1
