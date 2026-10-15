"""Выгрузка результатов в CSV и XLSX (openpyxl)."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import CATEGORIES, COLUMN_TITLES, ROW_COLUMNS, URGENCIES, ProcessedLead

_WIDTHS = {
    "id": 9, "source": 9, "received_at": 17, "category": 13, "urgency": 11, "name": 18, "phone": 15,
    "service": 22, "desired_date": 20, "budget": 11, "summary": 48, "subject": 26, "text": 60,
    "provider": 10, "prompt_version": 9, "elapsed_ms": 9, "error": 30,
}
_URGENCY_FILL = {
    "высокая": PatternFill("solid", fgColor="FFC7CE"),
    "средняя": PatternFill("solid", fgColor="FFEB9C"),
    "низкая": PatternFill("solid", fgColor="E2EFDA"),
}


def write_csv(results: list[ProcessedLead], path: str | Path) -> Path:
    """CSV в UTF-8 с BOM — такой файл Excel открывает с правильной кодировкой."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(ROW_COLUMNS))
        writer.writeheader()
        for result in results:
            writer.writerow(result.to_row())
    return path


def write_xlsx(results: list[ProcessedLead], path: str | Path) -> Path:
    """XLSX с листом «Заявки» (фильтр, закреплённая шапка, подсветка срочности) и листом «Сводка»."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    ws.append([COLUMN_TITLES[c] for c in ROW_COLUMNS])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    urgency_col = ROW_COLUMNS.index("urgency") + 1
    budget_col = ROW_COLUMNS.index("budget") + 1
    text_col = ROW_COLUMNS.index("text") + 1
    for result in results:
        row = result.to_row()
        ws.append([row[c] for c in ROW_COLUMNS])
        r = ws.max_row
        fill = _URGENCY_FILL.get(row["urgency"])
        if fill:
            ws.cell(row=r, column=urgency_col).fill = fill
        ws.cell(row=r, column=budget_col).number_format = "#,##0"
        ws.cell(row=r, column=text_col).alignment = Alignment(wrap_text=True, vertical="top")

    for i, column in enumerate(ROW_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = _WIDTHS.get(column, 14)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Сводка")
    summary.append(["Показатель", "Значение"])
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    ok = [r for r in results if r.ok]
    summary.append(["Всего заявок", len(results)])
    summary.append(["Обработано", len(ok)])
    summary.append(["Ошибок", len(results) - len(ok)])
    summary.append([])
    summary.append(["Категория", "Заявок"])
    categories = Counter(r.analysis.category for r in ok)
    for category in CATEGORIES:
        summary.append([category, categories.get(category, 0)])
    summary.append([])
    summary.append(["Срочность", "Заявок"])
    urgencies = Counter(r.analysis.urgency for r in ok)
    for urgency in URGENCIES:
        summary.append([urgency, urgencies.get(urgency, 0)])
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 12

    wb.save(path)
    return path
